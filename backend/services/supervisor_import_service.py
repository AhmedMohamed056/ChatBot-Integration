"""Staged supervisor + campaign ownership import from Excel."""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select

from db.base import get_session
from db.models import Campaign
from db.platform_models import ImportBatch, Supervisor
from db.repositories.platform_repository import SupervisorRepository
from phone_utils import normalize_phone
from services.runtime_settings import get_runtime_setting

REQUIRED_COLUMNS = {"campaign name", "supervisor name", "phone number"}
OPTIONAL_COLUMNS = {"status", "imported date"}


def _headers(row) -> dict[str, int]:
    normalised = {}
    for idx, cell in enumerate(row):
        if cell is None:
            continue
        key = str(cell).strip().lower()
        if key:
            normalised[key] = idx
    missing = REQUIRED_COLUMNS - set(normalised.keys())
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(sorted(missing))}")
    mapping = {col: normalised[col] for col in REQUIRED_COLUMNS}
    for opt in OPTIONAL_COLUMNS:
        if opt in normalised:
            mapping[opt] = normalised[opt]
    return mapping


def preview_supervisor_import() -> dict[str, Any]:
    path = Path(get_runtime_setting("campaign_file", ""))
    if not path.exists():
        raise FileNotFoundError("Campaign/supervisor file not configured.")
    wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    rows = list(wb.active.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Excel file is empty.")
    col = _headers([str(c or "") for c in rows[0]])
    preview = []
    errors = []
    seen_phones: dict[str, int] = {}
    for i, row in enumerate(rows[1:], start=2):
        try:
            campaign_name = str(row[col["campaign name"]] or "").strip()
            supervisor_name = str(row[col["supervisor name"]] or "").strip()
            phone_raw = str(row[col["phone number"]] or "").strip()
            phone = normalize_phone(phone_raw)
            status = (
                str(row[col["status"]]).strip().lower()
                if "status" in col and row[col["status"]] is not None
                else "active"
            )
            if not (campaign_name and supervisor_name and phone):
                raise ValueError("missing required values")
            if phone in seen_phones:
                errors.append(
                    {
                        "row": i,
                        "error": f"duplicate phone {phone} (also on row {seen_phones[phone]})",
                    }
                )
                continue
            seen_phones[phone] = i
            preview.append(
                {
                    "row": i,
                    "campaign_name": campaign_name,
                    "supervisor_name": supervisor_name,
                    "phone_number": phone,
                    "status": status,
                }
            )
        except Exception as exc:  # noqa: BLE001
            errors.append({"row": i, "error": str(exc)})
    return {"ok": True, "preview": preview, "errors": errors}


def activate_supervisor_import() -> dict[str, Any]:
    preview = preview_supervisor_import()
    if preview["errors"]:
        dupes = [e for e in preview["errors"] if "duplicate phone" in e.get("error", "")]
        if dupes:
            raise ValueError(
                "Phone conflicts in Excel: "
                + "; ".join(e["error"] for e in dupes[:5])
            )
    activated = 0
    imported_phones: set[str] = set()
    with get_session() as session:
        batch = ImportBatch(
            import_type="supervisor_campaign",
            status="activating",
            row_count=len(preview["preview"]),
            error_count=len(preview["errors"]),
        )
        session.add(batch)
        session.flush()
        sup_repo = SupervisorRepository(session)
        for row in preview["preview"]:
            phone = row["phone_number"]
            imported_phones.add(phone)
            supervisor = sup_repo.get_by_phone(phone)
            if supervisor is None:
                supervisor = sup_repo.create(
                    {
                        "phone_number": phone,
                        "display_name": row["supervisor_name"],
                        "is_active": row.get("status", "active") != "inactive",
                        "activated_at": datetime.now(timezone.utc),
                    }
                )
            else:
                supervisor.display_name = row["supervisor_name"]
                supervisor.is_active = row.get("status", "active") != "inactive"
                supervisor.deactivated_at = None
                supervisor.activated_at = datetime.now(timezone.utc)
            existing = session.scalar(
                select(Campaign).where(Campaign.campaign_name == row["campaign_name"])
            )
            if existing is None:
                session.add(
                    Campaign(
                        campaign_name=row["campaign_name"],
                        owner_supervisor_id=supervisor.id,
                        status="active",
                        description="",
                    )
                )
            else:
                existing.owner_supervisor_id = supervisor.id
                if existing.status == "deleted":
                    existing.status = "active"
            activated += 1

        for inactive in session.scalars(select(Supervisor)).all():
            if inactive.phone_number not in imported_phones and inactive.is_active:
                inactive.is_active = False
                inactive.deactivated_at = datetime.now(timezone.utc)

        batch.status = "active"
        batch.result = {"activated": activated, "imported_phones": len(imported_phones)}
    return {"ok": True, "activated": activated, "errors": preview["errors"]}
